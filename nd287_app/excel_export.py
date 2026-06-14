# -*- coding: utf-8 -*-
"""Excel(.xlsx)出力（表＋グラフ）

openpyxl が無い環境でも import は通り、書き出し時にだけ分かりやすい
エラーを出す（exe化時は同梱する想定）。

sheets の各要素:
    name    : シート名
    headers : 見出し行（リスト）
    rows    : データ行（リストのリスト）
    chart   : 省略可。ネイティブExcelグラフを表に基づいて作る:
              dict(type="bar"|"line", title, cat_col, val_cols,
                   anchor, x_title, y_title)
              cat_col / val_cols は headers の 0始まり列番号
    images  : 省略可。dict(path, anchor, width, height) のリスト（PNG等を貼る）
"""

try:
    import openpyxl
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    HAVE_OPENPYXL = True
except ImportError:  # pragma: no cover - 依存が無い環境向け
    HAVE_OPENPYXL = False


def _require():
    if not HAVE_OPENPYXL:
        raise RuntimeError(
            "Excel出力には openpyxl が必要です（pip install openpyxl）"
        )


def _style_header(ws, headers):
    if not headers:
        return
    fill = PatternFill("solid", fgColor="DDEBF7")
    font = Font(bold=True)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")


def _autosize(ws, headers, rows):
    widths = [len(str(h)) for h in headers] if headers else []
    for row in rows:
        for i, value in enumerate(row):
            if i >= len(widths):
                widths.append(0)
            widths[i] = max(widths[i], len(str(value)))
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(w + 2, 6), 40)


def _add_chart(ws, spec, headers, nrows):
    ctype = spec.get("type", "bar")
    chart = BarChart() if ctype == "bar" else LineChart()
    chart.title = spec.get("title", "")
    chart.height = spec.get("height", 9)
    chart.width = spec.get("width", 20)
    chart.y_axis.title = spec.get("y_title", "")
    chart.x_axis.title = spec.get("x_title", "")
    first, last = 2, nrows + 1
    for vc in spec["val_cols"]:
        ref = Reference(ws, min_col=vc + 1, min_row=1, max_row=last)
        chart.add_data(ref, titles_from_data=True)
    cats = Reference(ws, min_col=spec["cat_col"] + 1, min_row=first, max_row=last)
    chart.set_categories(cats)
    ws.add_chart(chart, spec.get("anchor", f"A{nrows + 4}"))


def _add_image(ws, img):
    picture = XLImage(img["path"])
    if img.get("width"):
        picture.width = img["width"]
    if img.get("height"):
        picture.height = img["height"]
    ws.add_image(picture, img.get("anchor", "A1"))


def write_xlsx(path, sheets):
    """シート定義のリストを1つの .xlsx に書き出す。"""
    _require()
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet in sheets:
        ws = wb.create_sheet(title=(sheet.get("name") or "データ")[:31])
        headers = sheet.get("headers") or []
        rows = sheet.get("rows") or []
        if headers:
            ws.append(list(headers))
        for row in rows:
            ws.append(list(row))
        _style_header(ws, headers)
        _autosize(ws, headers, rows)
        if headers:
            ws.freeze_panes = "A2"
        chart = sheet.get("chart")
        if chart and rows and chart.get("val_cols"):
            _add_chart(ws, chart, headers, len(rows))
        for img in sheet.get("images") or []:
            _add_image(ws, img)
    wb.save(str(path))
    return str(path)
