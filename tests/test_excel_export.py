# -*- coding: utf-8 -*-
import tempfile
import unittest
from pathlib import Path

from nd287_app import excel_export


@unittest.skipUnless(excel_export.HAVE_OPENPYXL, "openpyxl 未インストール")
class TestExcelExport(unittest.TestCase):
    def _write(self, sheets):
        import openpyxl
        tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        tmp.close()
        excel_export.write_xlsx(tmp.name, sheets)
        return openpyxl.load_workbook(tmp.name), tmp.name

    def test_table_written(self):
        headers = ["機番", "精度PP"]
        rows = [["10001", "12.30"], ["10002", "9.40"]]
        wb, path = self._write([dict(name="比較", headers=headers, rows=rows)])
        ws = wb["比較"]
        self.assertEqual(ws["A1"].value, "機番")
        self.assertEqual(ws["B2"].value, "12.30")
        self.assertEqual(ws.max_row, 3)
        Path(path).unlink()

    def test_native_chart_added(self):
        headers = ["機番", "精度PP"]
        rows = [["10001", "12.30"], ["10002", "9.40"]]
        chart = dict(type="bar", title="精度PP", cat_col=0, val_cols=[1],
                     x_title="機番", y_title="秒")
        wb, path = self._write([dict(name="比較", headers=headers, rows=rows, chart=chart)])
        self.assertEqual(len(wb["比較"]._charts), 1)
        Path(path).unlink()

    def test_multiple_sheets(self):
        sheets = [
            dict(name="指標", headers=["項目", "値"], rows=[["精度PP", "12.3"]]),
            dict(name="偏差", headers=["角度", "ホイールCW"], rows=[["0", "1.0"], ["30", "2.0"]],
                 chart=dict(type="line", title="偏差", cat_col=0, val_cols=[1])),
        ]
        wb, path = self._write(sheets)
        self.assertEqual(wb.sheetnames, ["指標", "偏差"])
        self.assertEqual(len(wb["偏差"]._charts), 1)
        Path(path).unlink()

    def test_no_chart_when_no_rows(self):
        chart = dict(type="bar", title="x", cat_col=0, val_cols=[1])
        wb, path = self._write([dict(name="空", headers=["a", "b"], rows=[], chart=chart)])
        self.assertEqual(len(wb["空"]._charts), 0)
        Path(path).unlink()


class TestExcelGuard(unittest.TestCase):
    def test_requires_openpyxl(self):
        if not excel_export.HAVE_OPENPYXL:
            with self.assertRaises(RuntimeError):
                excel_export.write_xlsx("x.xlsx", [])


if __name__ == "__main__":
    unittest.main()
